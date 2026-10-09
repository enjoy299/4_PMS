"""待拍板清单桥：确认列 ⬜ 项 <-> Excel 双向桥（md 是唯一事实源，xlsx 只是业务填写视图）。

为什么需要这个桥：
  确认列的「源」永远是 md 产物（AI 分析与 _lint_confirmations 都只消费 md）；
  但业务人员更习惯在 Excel 里筛选、批注、批量填写。
  本脚本只做格式搬运，不做任何判断：

    export -> 扫描 L1~L4 确认列的 ⬜ 项，生成《待拍板清单_*.xlsx》
              （业务只填 F 列「确认结论(业务填)」，附「填写说明」sheet）
    apply   -> 读业务填好的 xlsx，把结论原位回写到源 md 的对应行（⬜ -> ✅ <结论>），
              与业务手工在 md 上就地拍板完全等效

回写安全（fail-loud，绝不误写）：
  - 定位 = 来源文件 + 行号 + 原问句文本 三重校验；
  - 导出后产物重跑导致行漂移 / 行已被就地拍板的，跳过并列入「过期冲突」报告，退出码 1；
  - 只回写业务实际填了结论的行；空行、未改动文件一律不动。

用法：
  python scripts/confirmation_bridge.py export [模块名 ...]      # 不带模块 = 全部模块
  python scripts/confirmation_bridge.py apply [xlsx路径] [--dry-run]
                                              # 不带路径 = 取 output/05_export 下最新一份清单

控制台不打印 ⬜/✅ 原字符（GBK 终端安全），显示为 [未确认]/[已确认]。
"""
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))          # 直接 `python scripts/...` 运行时保证可导入 scripts 包

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import scripts._lint_confirmations as lint   # 复用分层发现逻辑（数字前缀容忍/最新批次/排除 archive）

PENDING_GLOB = "待拍板清单_*.xlsx"
_HEADERS = ["序号", "模块", "层级", "编号/上下文", "待确认问句",
            "确认结论(业务填)", "备注", "来源文件", "行号"]


def _norm(s) -> str:
    """比较用规范化：去全部空白（Excel 编辑、md 排版都容易引入空格差异）。"""
    return "".join(str(s).split())


def _fmt(text) -> str:
    """控制台安全文本：⬜/✅ 转文字（GBK 终端无法编码），超宽截断。"""
    t = str(text).replace("⬜", "[未确认]").replace("✅", "[已确认]")
    t = t.encode("gbk", errors="replace").decode("gbk")
    return t[:80] + ("…" if len(t) > 80 else "")


def _confirm_cell(line: str, conclusion: str) -> str:
    """行内原位替换：从首个 ⬜ 起、到单元格边界（下一个 | 或行尾）改为 ✅ <结论>。

    替换区含原收尾空格，统一补回一个空格保持「结论 |」排版（与手工就地拍板的样式一致）。
    """
    idx = line.find("⬜")
    end = line.find("|", idx)
    if end == -1:
        end = len(line)
    payload = conclusion if conclusion.startswith("✅") else f"✅ {conclusion}"
    if end < len(line) and not payload.endswith(" "):
        payload += " "
    return line[:idx] + payload + line[end:]


# ---------- 导出侧 ----------

def collect_pending(modules: list[str]) -> list[dict]:
    """复用 lint 分层发现逻辑，收集 L1~L4 全部 ⬜ 行（含系统级业务全景图）。

    返回字段：module / layer / qid（编号或最近标题）/ question / rel（相对或绝对路径）/ line。
    """
    items: list[dict] = []

    def scan_file(path: Path, module: str, lname: str):
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except OSError as e:
            print(f"  [WARN] 读取失败，跳过：{path}（{e}）")
            return
        last_heading = ""
        for ln, line in enumerate(text.splitlines(), 1):
            s = line.strip()
            if s.startswith("#"):
                last_heading = s.lstrip("#").strip()
            if not s.startswith("|") or "⬜" not in s:
                continue
            cells = [c.strip() for c in s.split("|")]
            cidx = next((i for i, c in enumerate(cells) if "⬜" in c), len(cells))
            question = cells[cidx].split("⬜", 1)[1].strip() if cidx < len(cells) else ""
            # 编号/上下文：⬜ 前第一个非空单元格；竖排字段行（首格是「确认结论(业务填)」）退化为最近标题
            pre = [c for c in cells[1:cidx] if c]
            qid = pre[0] if pre and "确认结论" not in pre[0] else ""
            if not qid:
                qid = last_heading or "（无编号）"
            try:
                rel = str(path.relative_to(ROOT))
            except ValueError:
                rel = str(path)                     # 测试/自定义输出目录在项目外时给绝对路径
            items.append({"module": module, "layer": lname, "qid": qid,
                          "question": question, "rel": rel, "line": ln})

    if lint.BUSINESS_DIR.is_dir():
        for p in sorted(lint.BUSINESS_DIR.glob("00-*-业务全景图.md")):   # 系统级全景图影响所有模块
            scan_file(p, "系统级", lint.L1)
    for m in modules:
        stems = lint._req_stems(m)
        for lname, files in ((lint.L1, lint._layer1_files(m)),
                             (lint.L2, lint._layer2_files(m, stems)),
                             (lint.L3, lint._layer3_files(m)),
                             (lint.L4, lint._layer4_files(m))):
            for p in files:
                scan_file(p, m, lname)
    return items


def cmd_export(argv: list[str]) -> int:
    modules = [lint._clean(a) for a in argv if lint._clean(a)] or lint._modules()
    if not modules:
        print("未发现任何模块（output/01_business、02_split_requirements 为空）")
        return 0
    items = collect_pending(modules)
    if not items:
        print("RESULT: PASS（确认列已全部拍板，无需导出待拍板清单）")
        return 0

    scope = "_".join(modules) if len(modules) <= 3 else "全部"
    out = lint.EXPORT_DIR / f"待拍板清单_{scope}_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    lint.EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "待拍板"
    ws.append(_HEADERS)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for i, it in enumerate(items, 1):
        ws.append([i, it["module"], it["layer"], it["qid"], it["question"], "", "",
                   it["rel"], it["line"]])
    for col, w in {"A": 6, "B": 12, "C": 12, "D": 30, "E": 58, "F": 44,
                   "G": 16, "H": 60, "I": 8}.items():
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"

    note = wb.create_sheet("填写说明")
    for n in [
        "本清单由框架导出：来源是各产物 md「确认结论(业务填)」列中仍未拍板（⬜）的条目。",
        "只需填写 F 列「确认结论(业务填)」：直接写结论文字（不要保留问句）；不确定的行留空即可。",
        "请勿修改「待确认问句 / 来源文件 / 行号」等列——回写脚本靠它们定位原文，改了会回写失败。",
        "填完后执行：python scripts/confirmation_bridge.py apply <本文件路径>（可先加 --dry-run 预演）。",
        "回写成功后，源 md 该行原位变为「✅ <结论>」，与手工在 md 上就地拍板完全等效。",
    ]:
        note.append([n])
    note.column_dimensions["A"].width = 110

    wb.save(out)
    print(f"待拍板清单已导出：{out}")
    print(f"共 {len(items)} 项未拍板，请业务填写 F 列后执行 apply 回流")
    for it in items[:10]:
        print(f"  [{it['layer']}] {it['qid']}  {_fmt(it['question'])}")
    if len(items) > 10:
        print(f"  … 其余 {len(items) - 10} 项见 xlsx")
    return 0


# ---------- 回流侧 ----------

def cmd_apply(argv: list[str]) -> int:
    path_arg = next((a for a in argv if not a.startswith("-")), None)
    dry_run = "--dry-run" in argv
    if path_arg:
        xlsx = Path(path_arg)
    else:
        cands = (sorted(lint.EXPORT_DIR.glob(PENDING_GLOB), key=lambda p: p.stat().st_mtime)
                 if lint.EXPORT_DIR.is_dir() else [])
        if not cands:
            print("未找到待拍板清单（output/05_export/待拍板清单_*.xlsx）；请先运行 export")
            return 1
        xlsx = cands[-1]
        print(f"未指定路径，取最新清单：{xlsx.name}")
    if not xlsx.exists():
        print(f"文件不存在：{xlsx}")
        return 1

    wb = load_workbook(xlsx, data_only=True)
    if "待拍板" not in wb.sheetnames:
        print(f"缺少「待拍板」工作表，不是本脚本导出的清单：{xlsx.name}")
        return 1
    ws = wb["待拍板"]
    header = [str(c.value or "").strip() for c in ws[1]]
    try:
        col = {name: header.index(name) + 1 for name in _HEADERS}
    except ValueError as e:
        print(f"表头列缺失{e}，清单可能被改动：{xlsx.name}")
        return 1

    edits: dict[str, list] = defaultdict(list)     # 源文件 -> [(行号, 问句, 结论, 序号)]
    empty = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        no = row[col["序号"] - 1]
        if no is None:
            continue                                # 业务在表尾粘贴出的空行
        question = str(row[col["待确认问句"] - 1] or "").strip()
        conclusion = str(row[col["确认结论(业务填)"] - 1] or "").strip()
        src = str(row[col["来源文件"] - 1] or "").strip()
        line_no = row[col["行号"] - 1]
        if not src or not line_no:
            continue
        if not conclusion:
            empty += 1
            continue
        p = Path(src)
        if not p.is_absolute():
            p = ROOT / src
        try:
            n = int(line_no)
        except (TypeError, ValueError):
            continue
        edits[str(p)].append((n, question, conclusion, no))

    if not edits:
        print(f"业务未填写任何结论（空行 {empty} 条），无可回写")
        return 0

    applied, stale = [], []
    for src, rows in edits.items():
        path = Path(src)
        try:
            # newline=""：不做换行转换，逐字节保留原文件行尾（LF 文件不能被改写成 CRLF，
            # 否则整文件 diff 污染）；splitlines(keepends=True) 才能拿到真实 eol
            with path.open(encoding="utf-8", newline="") as f:
                raw = f.read()
        except (OSError, ValueError) as e:
            for _, _, _, no in rows:
                stale.append((no, path, "-", f"读取失败：{e}"))
            continue
        lines = raw.splitlines(keepends=True)      # keepends：保留每行原有换行符（含 CRLF）
        changed = False
        for line_no, question, conclusion, no in rows:
            if line_no < 1 or line_no > len(lines):
                stale.append((no, path, line_no, "行号越界（产物可能已重跑）"))
                continue
            target = lines[line_no - 1]
            body = target.rstrip("\r\n")
            if "⬜" not in body:
                stale.append((no, path, line_no, "该行已无未确认标记（已拍板或行漂移）"))
                continue
            if question and _norm(question) not in _norm(body):
                stale.append((no, path, line_no, "原问句不匹配（行内容已变化）"))
                continue
            lines[line_no - 1] = _confirm_cell(body, conclusion) + target[len(body):]
            changed = True
            applied.append((no, path, line_no, conclusion))
        if changed and not dry_run:
            with path.open("w", encoding="utf-8", newline="") as f:   # 不转换换行，只改目标行
                f.write("".join(lines))

    verb = "预演" if dry_run else "回写"
    print(f"待拍板清单回流（{verb}）：{xlsx.name}")
    print(f"  成功{verb} {len(applied)} 条 ｜ 未填跳过 {empty} 条 ｜ 过期冲突 {len(stale)} 条")
    for no, path, line_no, conclusion in applied:
        payload = conclusion if conclusion.startswith("✅") else f"✅ {conclusion}"
        print(f"  [已{verb}] #{no} {path.name}:{line_no}  ->  {_fmt(payload)}")
    for no, path, line_no, reason in stale:
        print(f"  [过期] #{no} {path.name}:{line_no}  {reason}")
    if stale:
        print("  过期项请重新 export 最新清单再填（产物在导出后发生过重跑 / 该行已被就地拍板）")
        return 1
    if not dry_run:
        print("  回流完成：源 md 已更新；可运行 python scripts/_lint_confirmations.py 复查剩余未拍板项")
    return 0


def main(argv: list[str]) -> int:
    if not argv or argv[0] not in ("export", "apply"):
        print(__doc__)
        return 0 if not argv else 1
    if argv[0] == "export":
        return cmd_export(argv[1:])
    return cmd_apply(argv[1:])


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
