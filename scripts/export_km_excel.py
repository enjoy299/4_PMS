"""
导出知识管理模块的 测试点+测试用例 合并 Excel 文件。

⚠ 已弃用：推荐使用 `python scripts/export_module_excel.py 知识管理`
本脚本保留作为向后兼容入口。

输入（新目录约定）：
    output/03_design/test_points/知识管理/01-*.md
    output/03_design/test_cases/知识管理/01-*.md
输出：output/05_export/知识管理/07-知识管理-测试点+测试用例_{时间戳}.xlsx

用法：python scripts/export_km_excel.py
"""
import re
import sys
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).resolve().parent.parent
MODULE_NUM = "07"  # 7、知识管理 → 07
MODULE_NAME = "知识管理"
TP_DIR = ROOT / "output/03_design/test_points" / MODULE_NAME
CASE_DIR = ROOT / "output/03_design/test_cases" / MODULE_NAME
OUTPUT_DIR = ROOT / "output/05_export/知识管理"

# 需要合并的子模块文件（对应知识管理目录）
SUB_FILES = ["分类管理", "知识查询", "知识维护"]


def parse_test_points(md_text: str) -> list[dict]:
    """解析 8 列测试点 Markdown 表格。"""
    rows = []
    tp_headers = ["测试点编号", "模块", "需求来源", "测试维度", "测试点描述", "优先级", "是否依赖确认", "备注"]
    seen_ids = set()
    for line in md_text.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        if line.startswith("|---") or line.startswith("|:---"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) < 8:
            continue
        # 跳过表头行本身
        if cells[0] == "测试点编号":
            continue
        # 数据行必须以 TP- 开头
        if not cells[0].startswith("TP-"):
            continue
        if cells[0] in seen_ids:
            continue
        seen_ids.add(cells[0])
        rows.append({tp_headers[i]: cells[i] for i in range(8)})
    return rows


def parse_test_cases(md_text: str) -> list[dict]:
    """解析 13 字段竖排测试用例（### TC-XXX + 属性表格）。"""
    cases = []
    current = None
    case_fields = ["用例ID", "模块", "需求来源", "测试点ID", "用例标题", "优先级", "用例类型",
                   "前置条件", "测试数据", "操作步骤", "预期结果", "清理动作", "备注"]

    for line in md_text.splitlines():
        s = line.strip()
        if s.startswith("### TC-"):
            if current and current.get("用例ID"):
                cases.append(current)
            m = re.match(r"###\s+(TC-\S+)\s+(.*)", s)
            current = {k: "" for k in case_fields}
            if m:
                current["用例ID"] = m.group(1)
                current["用例标题"] = m.group(2).strip()
            continue
        if current is None:
            continue
        if not s.startswith("|"):
            continue
        if s.startswith("|---") or s.startswith("|:---"):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if len(cells) < 2:
            continue
        key, val = cells[0], cells[1]
        val = val.replace("<br>", "\n").replace("<BR>", "\n")
        if key == "字段" or key == "内容":  # 表头
            continue
        if key in current:
            current[key] = val
    if current and current.get("用例ID"):
        cases.append(current)
    return cases


def build_excel(tp_data: list[dict], case_data: list[dict], output_path: Path):
    """生成双 Sheet Excel（测试点 + 测试用例），带表头样式。"""
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    # 优先级颜色
    p0_fill = PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid")
    p1_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
    p2_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")

    def color_for(p):
        return {"P0": p0_fill, "P1": p1_fill, "P2": p2_fill}.get(p)

    # ===== Sheet 1: 测试点 =====
    ws1 = wb.active
    ws1.title = "测试点"
    tp_headers = ["测试点编号", "模块", "需求来源", "测试维度", "测试点描述", "优先级", "是否依赖确认", "备注"]
    for col, h in enumerate(tp_headers, 1):
        c = ws1.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = thin_border
    for r, row in enumerate(tp_data, 2):
        for col, h in enumerate(tp_headers, 1):
            c = ws1.cell(row=r, column=col, value=row.get(h, ""))
            c.alignment = cell_align
            c.border = thin_border
            if h == "优先级":
                fill = color_for(row.get(h, ""))
                if fill:
                    c.fill = fill
    for i, w in enumerate([14, 12, 15, 12, 45, 8, 12, 30], 1):
        ws1.column_dimensions[ws1.cell(row=1, column=i).column_letter].width = w
    ws1.freeze_panes = "A2"

    # ===== Sheet 2: 测试用例 =====
    ws2 = wb.create_sheet("测试用例")
    case_headers = ["用例ID", "模块", "需求来源", "测试点ID", "用例标题", "优先级", "用例类型",
                    "前置条件", "测试数据", "操作步骤", "预期结果", "清理动作", "备注"]
    for col, h in enumerate(case_headers, 1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = thin_border
    for r, row in enumerate(case_data, 2):
        for col, h in enumerate(case_headers, 1):
            c = ws2.cell(row=r, column=col, value=row.get(h, ""))
            c.alignment = cell_align
            c.border = thin_border
            if h == "优先级":
                fill = color_for(row.get(h, ""))
                if fill:
                    c.fill = fill
    for i, w in enumerate([14, 12, 15, 14, 40, 8, 10, 30, 25, 45, 40, 20, 20], 1):
        ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = w
    ws2.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"已生成：{output_path}")
    print(f"  测试点：{len(tp_data)} 条")
    print(f"  测试用例：{len(case_data)} 条")


def main():
    all_tp = []
    all_cases = []
    for sub in SUB_FILES:
        tp_file = TP_DIR / f"01-{sub}-测试点.md"
        case_file = CASE_DIR / f"01-{sub}-用例.md"
        if tp_file.exists():
            all_tp.extend(parse_test_points(tp_file.read_text(encoding="utf-8")))
            print(f"[读取] {tp_file.name}: {len(parse_test_points(tp_file.read_text(encoding='utf-8')))} 条测试点")
        else:
            print(f"[缺失] {tp_file}", file=sys.stderr)
        if case_file.exists():
            parsed_cases = parse_test_cases(case_file.read_text(encoding="utf-8"))
            all_cases.extend(parsed_cases)
            print(f"[读取] {case_file.name}: {len(parsed_cases)} 条用例")
        else:
            print(f"[缺失] {case_file}", file=sys.stderr)

    # 按 ID 排序（稳定输出）
    all_tp.sort(key=lambda x: x.get("测试点编号", ""))
    all_cases.sort(key=lambda x: x.get("用例ID", ""))

    ts = datetime.now().strftime("%Y%m%d_%H%M")
    output_path = OUTPUT_DIR / f"{MODULE_NUM}-{MODULE_NAME}-测试点+测试用例_{ts}.xlsx"
    build_excel(all_tp, all_cases, output_path)


if __name__ == "__main__":
    main()
