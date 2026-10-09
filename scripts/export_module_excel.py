"""
通用测试点+测试用例合并 Excel 导出脚本。
用法：
    python scripts/export_module_excel.py <模块名> [需求编号]
示例：
    python scripts/export_module_excel.py 工业APP 11
    python scripts/export_module_excel.py 知识管理 07
    python scripts/export_module_excel.py 项目管理 02

约定（新）：
- 输入按业务模块归类到子目录，文件名带时间戳：
    output/03_design/test_points/{模块名}/{文件 stem}-{类型}_{YYYYMMDD_HHMM}.md
    output/03_design/test_cases/{模块名}/{文件 stem}-用例_{YYYYMMDD_HHMM}.md
- 与 output/02_split_requirements/{模块名}/ 保持对齐。
- 同一 stem 存在多份时，默认取文件名排序最后（即时间最新）的一份；无时间戳的旧命名仍兼容。
- 若模块拆分为多个子文件（如知识管理下 01-分类管理/01-知识查询/01-知识维护），
  则通过 MODULE_CONFIG 配置要合并的 stem 列表；否则合并单个整合文件。
- 输出路径：output/05_export/{模块名}/{需求编号}-{模块名}-测试点+测试用例_{时间戳}.xlsx
"""
import re
import sys
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).resolve().parent.parent
TP_DIR = ROOT / "output/03_design/test_points"
CASE_DIR = ROOT / "output/03_design/test_cases"
EXPORT_ROOT = ROOT / "output/05_export"

# 各模块的子文件配置：{模块名: (需求编号, [子文件 stem 列表])}
# 若某模块只有一个整合文件，列表只包含该 stem
MODULE_CONFIG = {
    "工业APP": ("11", ["11-工业APP"]),
    "知识管理": ("07", ["01-分类管理", "01-知识查询", "01-知识维护"]),
    "项目管理": ("02", ["01-项目管理", "02-分类管理", "03-项目维护", "04-项目数据查询"]),
    "项目详情管理": ("03", ["03-项目详情管理"]),
}


def find_latest(dir: Path, stem: str, type_suffix: str) -> Path | None:
    """在 dir 中查找 {stem}-{type_suffix}.md 或 {stem}-{type_suffix}_*.md，返回时间最新的。

    排序规则：文件名排序（YYYYMMDD_HHMM 保证单调）取最后一个，同目录无时间戳的旧命名也会入选。
    """
    if not dir.exists():
        return None
    candidates = list(dir.glob(f"{stem}-{type_suffix}.md")) + list(dir.glob(f"{stem}-{type_suffix}_*.md"))
    if not candidates:
        return None
    return sorted(candidates)[-1]


def parse_test_points(md_text: str) -> list[dict]:
    """解析 8 列测试点 Markdown 表格。"""
    rows = []
    tp_headers = ["测试点编号", "模块", "需求来源", "测试维度", "测试点描述", "优先级", "是否依赖确认", "备注"]
    seen_ids = set()
    for line in md_text.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        if line.startswith("|---") or line.startswith("|:---"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) < 8:
            continue
        if cells[0] == "测试点编号":
            continue
        if not cells[0].startswith("TP-"):
            continue
        if cells[0] in seen_ids:
            continue
        seen_ids.add(cells[0])
        rows.append({tp_headers[i]: cells[i] for i in range(8)})
    return rows


def parse_test_cases(md_text: str) -> list[dict]:
    """解析 13 字段竖排测试用例（### TC-XXX + 属性表格）。"""
    cases = []
    current = None
    case_fields = ["用例ID", "模块", "需求来源", "测试点ID", "用例标题", "优先级", "用例类型",
                   "前置条件", "测试数据", "操作步骤", "预期结果", "清理动作", "备注"]

    for line in md_text.splitlines():
        s = line.strip()
        if s.startswith("### TC-"):
            if current and current.get("用例ID"):
                cases.append(current)
            m = re.match(r"###\s+(TC-\S+)\s+(.*)", s)
            current = {k: "" for k in case_fields}
            if m:
                current["用例ID"] = m.group(1)
                current["用例标题"] = m.group(2).strip()
            continue
        if current is None:
            continue
        if not s.startswith("|"):
            continue
        if s.startswith("|---") or s.startswith("|:---"):
            continue
        cells = [c.strip() for c in s.strip("|").split("|")]
        if len(cells) < 2:
            continue
        key, val = cells[0], cells[1]
        val = val.replace("<br>", "\n").replace("<BR>", "\n")
        if key in ("字段", "内容", "值"):
            continue
        if key in current:
            current[key] = val
    if current and current.get("用例ID"):
        cases.append(current)
    return cases


def build_excel(tp_data: list[dict], case_data: list[dict], output_path: Path):
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    p0_fill = PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid")
    p1_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")
    p2_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")

    def color_for(p):
        return {"P0": p0_fill, "P1": p1_fill, "P2": p2_fill}.get(p)

    ws1 = wb.active
    ws1.title = "测试点"
    tp_headers = ["测试点编号", "模块", "需求来源", "测试维度", "测试点描述", "优先级", "是否依赖确认", "备注"]
    for col, h in enumerate(tp_headers, 1):
        c = ws1.cell(row=1, column=col, value=h)
        c.font = header_font; c.fill = header_fill; c.alignment = header_align; c.border = thin_border
    for r, row in enumerate(tp_data, 2):
        for col, h in enumerate(tp_headers, 1):
            c = ws1.cell(row=r, column=col, value=row.get(h, ""))
            c.alignment = cell_align; c.border = thin_border
            if h == "优先级":
                fill = color_for(row.get(h, ""))
                if fill: c.fill = fill
    for i, w in enumerate([14, 14, 18, 12, 45, 8, 12, 30], 1):
        ws1.column_dimensions[ws1.cell(row=1, column=i).column_letter].width = w
    ws1.freeze_panes = "A2"

    ws2 = wb.create_sheet("测试用例")
    case_headers = ["用例ID", "模块", "需求来源", "测试点ID", "用例标题", "优先级", "用例类型",
                    "前置条件", "测试数据", "操作步骤", "预期结果", "清理动作", "备注"]
    for col, h in enumerate(case_headers, 1):
        c = ws2.cell(row=1, column=col, value=h)
        c.font = header_font; c.fill = header_fill; c.alignment = header_align; c.border = thin_border
    for r, row in enumerate(case_data, 2):
        for col, h in enumerate(case_headers, 1):
            c = ws2.cell(row=r, column=col, value=row.get(h, ""))
            c.alignment = cell_align; c.border = thin_border
            if h == "优先级":
                fill = color_for(row.get(h, ""))
                if fill: c.fill = fill
    for i, w in enumerate([14, 14, 18, 14, 40, 8, 12, 30, 25, 45, 40, 20, 20], 1):
        ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = w
    ws2.freeze_panes = "A2"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main(module_name: str, req_num: str | None = None):
    if module_name not in MODULE_CONFIG:
        print(f"未配置的模块：{module_name}，可选：{list(MODULE_CONFIG.keys())}", file=sys.stderr)
        sys.exit(1)
    default_num, stems = MODULE_CONFIG[module_name]
    num = req_num or default_num

    all_tp, all_cases = [], []
    seen_tp, seen_tc = set(), set()
    tp_module_dir = TP_DIR / module_name
    case_module_dir = CASE_DIR / module_name
    for stem in stems:
        tp_file = find_latest(tp_module_dir, stem, "测试点")
        case_file = find_latest(case_module_dir, stem, "用例")
        if tp_file and tp_file.exists():
            parsed = parse_test_points(tp_file.read_text(encoding="utf-8"))
            new_rows = [r for r in parsed if r["测试点编号"] not in seen_tp]
            for r in new_rows:
                seen_tp.add(r["测试点编号"])
            all_tp.extend(new_rows)
            print(f"[读取] {tp_file.name}: {len(new_rows)} 条测试点")
        else:
            print(f"[缺失] {tp_module_dir / (stem + '-测试点*.md')}", file=sys.stderr)
        if case_file and case_file.exists():
            parsed = parse_test_cases(case_file.read_text(encoding="utf-8"))
            new_rows = [r for r in parsed if r["用例ID"] not in seen_tc]
            for r in new_rows:
                seen_tc.add(r["用例ID"])
            all_cases.extend(new_rows)
            print(f"[读取] {case_file.name}: {len(new_rows)} 条用例")
        else:
            print(f"[缺失] {case_module_dir / (stem + '-用例*.md')}", file=sys.stderr)

    all_tp.sort(key=lambda x: x.get("测试点编号", ""))
    all_cases.sort(key=lambda x: x.get("用例ID", ""))

    ts = datetime.now().strftime("%Y%m%d_%H%M")
    output_path = EXPORT_ROOT / module_name / f"{num}-{module_name}-测试点+测试用例_{ts}.xlsx"
    build_excel(all_tp, all_cases, output_path)
    print(f"\n已生成：{output_path}")
    print(f"  测试点：{len(all_tp)} 条")
    print(f"  测试用例：{len(all_cases)} 条")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print("用法：python scripts/export_module_excel.py <模块名> [需求编号]", file=sys.stderr)
        sys.exit(1)
    mname = sys.argv[1]
    mnum = sys.argv[2] if len(sys.argv) > 2 else None
    main(mname, mnum)
