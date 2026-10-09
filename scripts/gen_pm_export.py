"""
生成项目管理模块的 测试点+测试用例 合并 Excel 文件。
格式参考：11-工业APP-测试点+测试用例_20260903_1041.xlsx
两个 Sheet：测试点 + 测试用例
"""
import re
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(r"E:\Qodertest\air_sys_test")
TP_FILE = ROOT / "output/03_design/test_points/01-项目管理-测试点.md"
CASE_FILE = ROOT / "output/03_design/test_cases/01-项目管理-用例.md"
OUTPUT_DIR = ROOT / "output/05_export/项目管理"

# 模块编号映射（从需求文件名提取）
MODULE_NUM = "02"  # 2、项目管理 → 02


def parse_test_points(md_text: str) -> list[dict]:
    """从 Markdown 解析测试点表格。"""
    rows = []
    current_module = ""
    # 章节名 → 模块名映射
    section_map = {
        "项目概述": "项目概述",
        "分类管理": "分类管理",
        "项目维护": "项目维护",
        "项目数据查询": "项目数据查询",
    }
    for line in md_text.splitlines():
        line = line.strip()
        # 识别章节标题
        if line.startswith("## "):
            section = line.lstrip("# ").replace("一、", "").replace("二、", "").replace("三、", "").replace("四、", "").strip()
            current_module = section_map.get(section, section)
            continue
        # 解析表格数据行
        if line.startswith("|") and "TP-" in line:
            cells = [c.strip() for c in line.split("|")[1:-1]]
            if len(cells) >= 3 and cells[0].startswith("TP-"):
                rows.append({
                    "测试点编号": cells[0],
                    "模块": current_module,
                    "需求来源": "-",
                    "测试维度": "正向功能",
                    "测试点描述": cells[1],
                    "优先级": cells[2],
                    "是否依赖确认": "否",
                    "备注": "-",
                })
    return rows


def parse_test_cases(md_text: str) -> list[dict]:
    """从 Markdown 解析测试用例。"""
    cases = []
    current_case = None

    for line in md_text.splitlines():
        line_stripped = line.strip()
        # 新用例标题
        if line_stripped.startswith("### TC-"):
            if current_case:
                cases.append(current_case)
            # 提取 ID 和标题
            m = re.match(r"### (TC-\S+)\s+(.*)", line_stripped)
            current_case = {
                "用例ID": m.group(1) if m else "",
                "用例标题": m.group(2) if m else "",
                "模块": "项目管理",
                "需求来源": "-",
                "测试点ID": "-",
                "用例类型": "正向",
                "前置条件": "",
                "测试数据": "-",
                "操作步骤": "",
                "预期结果": "",
                "清理动作": "-",
                "备注": "-",
                "优先级": "",
            }
            continue
        if current_case and line_stripped.startswith("|") and not line_stripped.startswith("|---"):
            cells = [c.strip() for c in line_stripped.split("|")[1:-1]]
            if len(cells) >= 2:
                key = cells[0].strip()
                val = cells[1].strip()
                # 将 <br> 替换为换行符
                val = val.replace("<br>", "\n")
                if key == "用例 ID":
                    current_case["用例ID"] = val
                elif key == "模块":
                    current_case["模块"] = val
                elif key == "子模块":
                    current_case["模块"] = val  # 用子模块作为模块值
                elif key == "优先级":
                    current_case["优先级"] = val
                elif key == "前置条件":
                    current_case["前置条件"] = val
                elif key == "测试步骤":
                    current_case["操作步骤"] = val
                elif key == "预期结果":
                    current_case["预期结果"] = val
                # 推断用例类型
                title = current_case.get("用例标题", "")
                if any(kw in title for kw in ["不", "禁止", "无法", "缺少", "不足", "不支持"]):
                    current_case["用例类型"] = "异常"
                elif any(kw in title for kw in ["禁用", "不可"]):
                    current_case["用例类型"] = "状态"

    if current_case:
        cases.append(current_case)
    return cases


def build_excel(tp_data: list[dict], case_data: list[dict], output_path: Path):
    """生成合并 Excel 文件。"""
    wb = Workbook()

    # ===== 样式定义 =====
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell_align = Alignment(vertical="top", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # ===== Sheet 1: 测试点 =====
    ws1 = wb.active
    ws1.title = "测试点"
    tp_headers = ["测试点编号", "模块", "需求来源", "测试维度", "测试点描述", "优先级", "是否依赖确认", "备注"]
    for col, h in enumerate(tp_headers, 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    for r, row_data in enumerate(tp_data, 2):
        for col, h in enumerate(tp_headers, 1):
            cell = ws1.cell(row=r, column=col, value=row_data.get(h, ""))
            cell.alignment = cell_align
            cell.border = thin_border
    # 列宽
    tp_widths = [14, 12, 10, 10, 35, 8, 12, 15]
    for i, w in enumerate(tp_widths, 1):
        ws1.column_dimensions[ws1.cell(row=1, column=i).column_letter].width = w
    ws1.freeze_panes = "A2"

    # ===== Sheet 2: 测试用例 =====
    ws2 = wb.create_sheet("测试用例")
    case_headers = ["用例ID", "模块", "需求来源", "测试点ID", "用例标题", "优先级", "用例类型",
                    "前置条件", "测试数据", "操作步骤", "预期结果", "清理动作", "备注"]
    for col, h in enumerate(case_headers, 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    for r, row_data in enumerate(case_data, 2):
        for col, h in enumerate(case_headers, 1):
            cell = ws2.cell(row=r, column=col, value=row_data.get(h, ""))
            cell.alignment = cell_align
            cell.border = thin_border
    # 列宽
    case_widths = [14, 12, 10, 12, 30, 8, 8, 35, 15, 35, 30, 15, 10]
    for i, w in enumerate(case_widths, 1):
        ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = w
    ws2.freeze_panes = "A2"

    # 保存
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f"已生成：{output_path}")
    print(f"  测试点：{len(tp_data)} 条")
    print(f"  测试用例：{len(case_data)} 条")


if __name__ == "__main__":
    tp_text = TP_FILE.read_text(encoding="utf-8")
    case_text = CASE_FILE.read_text(encoding="utf-8")

    tp_data = parse_test_points(tp_text)
    case_data = parse_test_cases(case_text)

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = OUTPUT_DIR / f"{MODULE_NUM}-项目管理-测试点+测试用例_{ts}.xlsx"
    build_excel(tp_data, case_data, output_path)
