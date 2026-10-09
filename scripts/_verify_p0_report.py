# -*- coding: utf-8 -*-
"""验证 P0 报告结构。"""
import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from openpyxl import load_workbook
from pathlib import Path

p = sorted(Path('output/06_reports/excel').glob('20260907_14*_知识管理_P0.xlsx'))[-1]
wb = load_workbook(p)
ws = wb.active
print(f'File: {p.name}')
print(f'Sheet: {ws.title}')
print(f'Dimensions: {ws.max_row} rows x {ws.max_column} cols')
print(f'A1: {ws["A1"].value}')
print(f'A2: {ws["A2"].value[:110] if ws["A2"].value else None}')
print('Header (row 4):')
for c in range(1, ws.max_column + 1):
    v = ws.cell(row=4, column=c).value
    print(f'  col {chr(64+c)}: {v}')
print()
print('Data rows:')
for r in range(5, ws.max_row + 1):
    cid = ws.cell(row=r, column=1).value
    st = ws.cell(row=r, column=6).value
    name = (ws.cell(row=r, column=2).value or '')[:32]
    ad = (ws.cell(row=r, column=14).value or '')[:80]
    print(f'  {cid:<12s} [{st:<5s}] {name:<32s} | {ad}')
