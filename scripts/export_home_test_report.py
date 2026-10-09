"""
首页 · 标准 Excel 测试报告生成器（test-report-format：14 字段 / Test Results Sheet）

数据来源（整合 06_reports/首页 的文档）：
- 测试执行记录：提供每条用例的 状态 + 实测结果（断言详情）
- 用例 md：提供 前置条件 / 操作步骤 / 预期结果
- META：模块 / 用例类型 / 需求来源 / 测试点
输出：output/06_reports/excel/{时间戳}_首页.xlsx（不建子目录，符合规范）
"""
import sys
import re
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).resolve().parent.parent
# 可通过 argv[1] 指定执行记录，argv[2] 指定输出文件名后缀（默认 20260907 记录）
EXEC_MD = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "output/06_reports/首页/20260907_测试执行记录.md"
OUT_SUFFIX = sys.argv[2] if len(sys.argv) > 2 else "首页"
CASES_MD = ROOT / "output/03_design/test_cases/首页/01-首页-用例_20260907_1900.md"
OUT_DIR = ROOT / "output/06_reports/excel"
OUT_DIR.mkdir(parents=True, exist_ok=True)

# 复用与用例导出一致的元数据
META = {
    "HOME-001": ("首页整体", "正向", "R0-1"), "HOME-002": ("首页整体", "权限", "R0-2"),
    "HOME-003": ("首页整体", "异常", "GAP-001"), "HOME-004": ("首页整体", "异常", "GAP-006"),
    "HOME-005": ("首页整体", "权限", "R0-2"),
    "TASK-001": ("我的任务", "正向", "R1-1"), "TASK-002": ("我的任务", "边界", "R1-1"),
    "TASK-003": ("我的任务", "异常", "GAP-E1"), "TASK-004": ("我的任务", "正向", "R1-3"),
    "REVT-001": ("我的待办审核", "正向", "R2-1"), "REVT-002": ("我的待办审核", "正向", "R2-3"),
    "REVT-003": ("我的待办审核", "异常", "GAP-E1"),
    "REVI-001": ("我发起的审核", "正向", "R3-1"), "REVI-002": ("我发起的审核", "正向", "R3-3"),
    "REVI-003": ("我发起的审核", "异常", "GAP-E1"),
    "PROJ-001": ("我的项目", "正向", "R4-1"), "PROJ-002": ("我的项目", "正向", "R4-3"),
    "PROJ-003": ("我的项目", "异常", "GAP-E1"),
    "STAT-001": ("项目概览", "正向", "R5-1"), "STAT-002": ("项目概览", "数据一致性", "R5-2"),
    "STAT-003": ("项目概览", "边界", "GAP-E4"),
    "TOOL-001": ("快捷工具栏", "配置", "GAP-017"),
}
HEADERS = ["Case ID", "Name", "Module", "Type", "Priority", "Status", "Duration",
           "Precondition", "Steps", "Expected", "Test Data", "Req Source", "Test Point", "Assertion Details"]
WIDTHS = [14, 30, 12, 10, 8, 9, 10, 30, 35, 30, 16, 12, 12, 45]

HDR_FONT = Font(bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", start_color="2E75B6", end_color="2E75B6")
THIN = Border(*[Side("thin")] * 4)
WRAP = Alignment(wrap_text=True, vertical="top")
FILL = {"pass": PatternFill("solid", start_color="E2EFDA", end_color="E2EFDA"),
        "fail": PatternFill("solid", start_color="FCE4EC", end_color="FCE4EC"),
        "block": PatternFill("solid", start_color="FFF3E0", end_color="FFF3E0"),
        "skip": PatternFill("solid", start_color="E7E6E6", end_color="E7E6E6")}


def clean(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    t = re.sub(r"<br\s*/?>", " ", t)
    return t.strip()


def norm_status(cell):
    u = cell.upper()
    if "PASS" in u or "✅" in cell:
        return "pass"
    if "FAIL" in u or "❌" in cell:
        return "fail"
    if "BLOCK" in u or "🚫" in cell:
        return "block"
    if "SKIP" in u or "⏭" in cell:
        return "skip"
    return "pass"


def parse_exec(md):
    """解析执行记录总表 -> {id: {priority,title,status,actual}}"""
    out = {}
    for ln in md.splitlines():
        m = re.match(r"^\|\s*(TC-[A-Z]+-\d+)\s*\|\s*(P\d)\s*\|\s*(.+?)\s*\|\s*(.+?)\s*\|\s*(.+?)\s*\|", ln)
        if m:
            cid, pri, title, res, actual = [clean(x) for x in m.groups()]
            out[cid] = {"priority": pri, "title": title,
                        "status": norm_status(res), "raw_status": res, "actual": actual}
    return out


def parse_cases(md):
    """解析用例 md -> {id: {precond,steps,expected}}"""
    out = {}
    for block in re.split(r"\n###\s+", "\n" + md)[1:]:
        lines = block.splitlines()
        title = clean(lines[0])
        cid = re.match(r"^(TC-[A-Z]+-\d+)", title)
        if not cid:
            continue
        cid = cid.group(1)
        f = {}
        for ln in lines:
            mm = re.match(r"^\|\s*(.+?)\s*\|\s*(.+?)\s*\|\s*$", ln)
            if mm:
                k, v = clean(mm.group(1)), clean(mm.group(2))
                if k != "项目" and not set(k) <= set("-: "):
                    f[k] = v
        out[cid] = {"precond": f.get("前置条件", ""), "steps": f.get("测试步骤", ""),
                    "expected": f.get("预期结果", "")}
    return out


def main():
    exec_rows = parse_exec(EXEC_MD.read_text(encoding="utf-8"))
    case_rows = parse_cases(CASES_MD.read_text(encoding="utf-8"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Results"
    ts_human = datetime.now().strftime("%Y-%m-%d %H:%M")
    n = len(exec_rows)
    npass = sum(1 for v in exec_rows.values() if v["status"] == "pass")
    rate = f"{npass/n*100:.0f}%" if n else "N/A"

    ws.merge_cells("A1:N1")
    ws.cell(1, 1, "首页 - Test Report").font = Font(bold=True, size=14)
    ws.merge_cells("A2:N2")
    ws.cell(2, 1, f"Time: {ts_human}   Pass Rate: {rate} ({npass}/{n})   "
                  f"PASS {npass} / BLOCK {sum(1 for v in exec_rows.values() if v['status']=='block')} / "
                  f"SKIP {sum(1 for v in exec_rows.values() if v['status']=='skip')} / "
                  f"FAIL {sum(1 for v in exec_rows.values() if v['status']=='fail')}")

    for col, h in enumerate(HEADERS, 1):
        c = ws.cell(4, col, h); c.font = HDR_FONT; c.fill = HDR_FILL; c.border = THIN
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # 保持用例顺序：按 META 定义顺序，再补其余
    order = [f"TC-{k}" for k in META.keys()]
    ids = [i for i in order if i in exec_rows] + [i for i in exec_rows if i not in order]

    r = 5
    for cid in ids:
        k = cid.split("-", 1)[1]
        mod, ctype, req = META.get(k, ("首页", "正向", "-"))
        e = exec_rows[cid]
        cs = case_rows.get(cid, {})
        tp = "TP-" + k
        assertion = f"{e['raw_status']}：{e['actual']}"
        if e["status"] == "fail":
            assertion += " [Error 见缺陷清单]"
        row = [cid, e["title"], mod, ctype, e["priority"], e["status"], "-",
               cs.get("precond", "-"), cs.get("steps", "-"), cs.get("expected", "-"),
               "账号/口令见仓库根 .env", req, tp, assertion]
        for col, val in enumerate(row, 1):
            c = ws.cell(r, col, val); c.border = THIN; c.alignment = WRAP
            if col == 6:
                c.fill = FILL.get(e["status"], FILL["pass"])
        r += 1

    for col, w in enumerate(WIDTHS, 1):
        ws.column_dimensions[ws.cell(4, col).column_letter].width = w
    ws.freeze_panes = "A5"

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out = OUT_DIR / f"{ts}_{OUT_SUFFIX}.xlsx"
    wb.save(out)
    print(f"已生成测试报告：{out}")
    print(f"用例 {n} 条：PASS {npass} / BLOCK {sum(1 for v in exec_rows.values() if v['status']=='block')} / "
          f"SKIP {sum(1 for v in exec_rows.values() if v['status']=='skip')} / "
          f"FAIL {sum(1 for v in exec_rows.values() if v['status']=='fail')}")


if __name__ == "__main__":
    main()
