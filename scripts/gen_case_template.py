"""
生成测试用例 Excel 模板文件。
模板包含 13 列标准字段，优先级带颜色区分（P0=红, P1=橙, P2=绿）。
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

TEMPLATE_DIR = Path(__file__).resolve().parent.parent / "templates"


def create_case_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "测试用例"

    # ── 样式 ──
    hdr_font = Font(bold=True, color="FFFFFF", size=10, name="微软雅黑")
    hdr_fill = PatternFill(start_color="2E75B6", end_color="2E75B6", fill_type="solid")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cell_align = Alignment(vertical="top", wrap_text=True)
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    # 优先级颜色
    p0_fill = PatternFill(start_color="FF4444", end_color="FF4444", fill_type="solid")  # 红
    p1_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")  # 橙
    p2_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")  # 绿
    p_font = Font(bold=True, color="FFFFFF", size=10)

    # ─ 13 列表头 ──
    headers = [
        "用例ID", "模块", "需求来源", "测试点ID", "用例标题",
        "优先级", "用例类型", "前置条件", "测试数据",
        "操作步骤", "预期结果", "清理动作", "备注",
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font, c.fill, c.alignment, c.border = hdr_font, hdr_fill, hdr_align, border

    # ── 示例行（P0） ──
    sample_p0 = [
        "TC-XXX-001", "示例模块", "REQ-01", "TP-XXX-001",
        "验证XXX功能正常", "P0", "正向",
        "已登录系统，进入XXX页面", "无",
        "1. 点击XXX按钮\n2. 输入XXX\n3. 点击保存",
        "XXX操作成功，页面显示正确结果",
        "删除创建的测试数据", "-",
    ]
    for col, val in enumerate(sample_p0, 1):
        c = ws.cell(row=2, column=col, value=val)
        c.alignment, c.border = cell_align, border
        if col == 6:  # 优先级列
            c.fill, c.font = p0_fill, p_font

    # ── 示例行（P1） ──
    sample_p1 = list(sample_p0)
    sample_p1[0] = "TC-XXX-002"
    sample_p1[4] = "验证XXX异常输入提示"
    sample_p1[5] = "P1"
    sample_p1[6] = "异常"
    sample_p1[7] = "已登录系统"
    sample_p1[8] = "输入非法值"
    sample_p1[9] = "1. 输入非法值\n2. 点击保存"
    sample_p1[10] = "提示错误信息"
    for col, val in enumerate(sample_p1, 1):
        c = ws.cell(row=3, column=col, value=val)
        c.alignment, c.border = cell_align, border
        if col == 6:
            c.fill, c.font = p1_fill, p_font

    # ── 示例行（P2） ──
    sample_p2 = list(sample_p0)
    sample_p2[0] = "TC-XXX-003"
    sample_p2[4] = "验证XXX边界值处理"
    sample_p2[5] = "P2"
    sample_p2[6] = "边界"
    for col, val in enumerate(sample_p2, 1):
        c = ws.cell(row=4, column=col, value=val)
        c.alignment, c.border = cell_align, border
        if col == 6:
            c.fill, c.font = p2_fill, p_font

    # ── 列宽 ──
    widths = [14, 12, 10, 12, 30, 8, 8, 35, 15, 35, 30, 15, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M1"

    # ── 保存 ──
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    path = TEMPLATE_DIR / "test_case_template.xlsx"
    wb.save(path)
    print(f"模板已生成：{path}")
    return path


if __name__ == "__main__":
    create_case_template()
