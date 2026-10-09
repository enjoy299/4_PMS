"""
将首页的「测试点 + 测试用例」Markdown 解析为符合 test-doc-format 标准的双 Sheet Excel。

- 直接解析：output/03_design/test_cases/首页/*.md 与 test_points/首页/*.md
- Sheet「测试点」8 列、Sheet「测试用例」13 列，样式沿用项目蓝底表头规范
- 输出：output/05_export/首页/{编号}-首页-测试点+测试用例_{ts}.xlsx
        并复制一份到 output/06_reports/excel/ 便于查看
"""
import re
import shutil
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

ROOT = Path(__file__).resolve().parent.parent
CASES_MD = ROOT / "output/03_design/test_cases/首页/01-首页-用例_20260907_1900.md"
POINTS_MD = ROOT / "output/03_design/test_points/首页/01-首页-测试点_20260907_1900.md"
OUT_DIR = ROOT / "output/05_export/首页"
COPY_DIR = ROOT / "output/06_reports/excel"
OUT_DIR.mkdir(parents=True, exist_ok=True)
COPY_DIR.mkdir(parents=True, exist_ok=True)

# ---- 样式 ----
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
HDR_FILL = PatternFill("solid", start_color="2E75B6", end_color="2E75B6")
HDR_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
CELL_ALIGN = Alignment(vertical="top", wrap_text=True)
THIN = Border(*[Side("thin")] * 4)

TP_HEADERS = ["测试点编号", "模块", "需求来源", "测试维度", "测试点描述", "优先级", "是否依赖确认", "备注"]
TP_WIDTHS = [14, 12, 10, 10, 35, 8, 12, 15]
TC_HEADERS = ["用例ID", "模块", "需求来源", "测试点ID", "用例标题", "优先级", "用例类型",
              "前置条件", "测试数据", "操作步骤", "预期结果", "清理动作", "备注"]
TC_WIDTHS = [14, 12, 10, 12, 30, 8, 8, 35, 15, 35, 30, 15, 10]

# ---- 按用例 ID 补充的元数据（需求来源 / 测试维度 / 用例类型 / 依赖确认 / 备注）----
META = {
    "HOME-001": ("R0-1", "正向功能", "正向", "否", "核心流程"),
    "HOME-002": ("R0-2", "权限控制", "权限", "否", "越权隔离核心"),
    "HOME-003": ("GAP-001", "异常输入", "异常", "否", "鉴权守卫"),
    "HOME-004": ("GAP-006", "异常输入", "异常", "是", "降级细则待确认"),
    "HOME-005": ("R0-2", "权限控制", "权限", "否", "双账号不串号"),
    "TASK-001": ("R1-1", "正向功能", "正向", "否", "-"),
    "TASK-002": ("R1-1", "边界值", "边界", "是", "排序口径=更新时间倒序待确认"),
    "TASK-003": ("GAP-E1", "异常输入", "异常", "是", "空态样式待确认"),
    "TASK-004": ("R1-3", "正向功能", "正向", "否", "-"),
    "REVT-001": ("R2-1", "正向功能", "正向", "否", "-"),
    "REVT-002": ("R2-3", "正向功能", "正向", "否", "-"),
    "REVT-003": ("GAP-E1", "异常输入", "异常", "是", "空态样式待确认"),
    "REVI-001": ("R3-1", "正向功能", "正向", "否", "-"),
    "REVI-002": ("R3-3", "正向功能", "正向", "否", "-"),
    "REVI-003": ("GAP-E1", "异常输入", "异常", "是", "空态样式待确认"),
    "PROJ-001": ("R4-1", "正向功能", "正向", "否", "-"),
    "PROJ-002": ("R4-3", "正向功能", "正向", "否", "-"),
    "PROJ-003": ("GAP-E1", "异常输入", "异常", "是", "空态样式待确认"),
    "STAT-001": ("R5-1", "正向功能", "正向", "是", "统计口径待确认"),
    "STAT-002": ("R5-2", "状态流转", "数据一致性", "否", "概览全量≠展示10"),
    "STAT-003": ("GAP-E4", "边界值", "边界", "是", "需零项目账号"),
    "TOOL-001": ("GAP-017", "正向功能", "配置", "是", "黑盒/无功能规范"),
}
SUBMOD = {"HOME": "首页整体", "TASK": "我的任务", "REVT": "我的待办审核",
          "REVI": "我发起的审核", "PROJ": "我的项目", "STAT": "项目概览", "TOOL": "快捷工具栏"}


def clean(t: str) -> str:
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    t = re.sub(r"<br\s*/?>", " ", t)
    return t.strip()


def key_of(cid: str) -> str:
    """TC-HOME-001 -> HOME-001；TP-TASK-002 -> TASK-002"""
    parts = cid.split("-")
    return "-".join(parts[1:]) if len(parts) >= 3 else cid


def submod_of(cid: str) -> str:
    parts = cid.split("-")
    return SUBMOD.get(parts[1] if len(parts) >= 2 else "", "首页")


def parse_cases(md: str):
    """解析竖排 key-value 表格的用例块，返回 list[dict]。"""
    cases = []
    for block in re.split(r"\n###\s+", "\n" + md)[1:]:
        lines = block.splitlines()
        raw_title = clean(lines[0])
        # 去除标题前缀的用例 ID（如 "TC-HOME-001 登录后…" -> "登录后…"）
        title = re.sub(r"^TC-[A-Z]+-\d+\s*", "", raw_title).strip()
        fields = {}
        for ln in lines:
            m = re.match(r"^\|\s*(.+?)\s*\|\s*(.+?)\s*\|\s*$", ln)
            if m:
                k, v = clean(m.group(1)), clean(m.group(2))
                if k in ("项目",) or set(k) <= set("-: "):
                    continue
                fields[k] = v
        cid = fields.get("用例 ID") or fields.get("用例ID") or title.split()[0]
        if not cid.startswith("TC"):
            continue
        cases.append({
            "id": cid, "title": title,
            "priority": fields.get("优先级", ""),
            "precond": fields.get("前置条件", ""),
            "steps": fields.get("测试步骤", ""),
            "expected": fields.get("预期结果", ""),
            "cleanup": fields.get("清理动作", "-"),
        })
    return cases


def parse_points(md: str):
    """解析测试点表格，返回 list[dict]（编号/描述/优先级）。"""
    pts = []
    for ln in md.splitlines():
        m = re.match(r"^\|\s*(TP-[A-Z]+-\d+)\s*\|\s*(.+?)\s*\|\s*(P\d)\s*\|\s*$", ln)
        if m:
            pts.append({"id": m.group(1), "desc": clean(m.group(2)), "priority": m.group(3)})
    return pts


def write_sheet(ws, headers, widths, rows):
    for col, h in enumerate(headers, 1):
        c = ws.cell(1, col, h); c.font = HDR_FONT; c.fill = HDR_FILL; c.alignment = HDR_ALIGN; c.border = THIN
    for i, row in enumerate(rows, 2):
        for col, val in enumerate(row, 1):
            c = ws.cell(i, col, val); c.alignment = CELL_ALIGN; c.border = THIN
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, col).column_letter].width = w
    ws.freeze_panes = "A2"


def main():
    cases = parse_cases(CASES_MD.read_text(encoding="utf-8"))
    points = parse_points(POINTS_MD.read_text(encoding="utf-8"))

    # ---- 测试点 Sheet ----
    tp_rows = []
    for p in points:
        k = key_of(p["id"])
        req, dim, _ctype, dep, note = META.get(k, ("-", "正向功能", "-", "否", "-"))
        tp_rows.append([p["id"], submod_of(p["id"]), req, dim, p["desc"], p["priority"], dep, note])

    # ---- 测试用例 Sheet ----
    tc_rows = []
    for c in cases:
        k = key_of(c["id"])
        req, _dim, ctype, _dep, note = META.get(k, ("-", "-", "正向", "否", "-"))
        tp_id = "TP-" + "-".join(c["id"].split("-")[1:])
        title = c["title"] if title_has_verb(c["title"]) else "验证" + c["title"]
        tc_rows.append([c["id"], submod_of(c["id"]), req, tp_id, title, c["priority"], ctype,
                        c["precond"] or "-", "账号/口令见仓库根 .env", c["steps"] or "-",
                        c["expected"] or "-", c["cleanup"] or "-", note])

    wb = Workbook()
    ws1 = wb.active; ws1.title = "测试点"
    write_sheet(ws1, TP_HEADERS, TP_WIDTHS, tp_rows)
    ws2 = wb.create_sheet("测试用例")
    write_sheet(ws2, TC_HEADERS, TC_WIDTHS, tc_rows)

    ts = datetime.now().strftime("%Y%m%d_%H%M")
    fname = f"01-首页-测试点+测试用例_{ts}.xlsx"
    out = OUT_DIR / fname
    wb.save(out)
    shutil.copy(out, COPY_DIR / fname)
    print(f"已生成：{out}")
    print(f"副本：  {COPY_DIR / fname}")
    print(f"测试点 {len(tp_rows)} 条 / 测试用例 {len(tc_rows)} 条")


def title_has_verb(t: str) -> bool:
    return t.startswith("验证")


if __name__ == "__main__":
    main()
