"""待拍板清单桥（confirmation_bridge）回归测试。

覆盖 export 导出（⬜ 收集/✅ 排除/竖排字段行/系统标题上下文）与 apply 回写
（原位 ✅ 替换/CRLF 保留/空行跳过/行漂移过期保护/--dry-run 预演）。
"""
from pathlib import Path

import pytest
from openpyxl import load_workbook

import scripts._lint_confirmations as lint
import scripts.confirmation_bridge as bridge

_TS = "20260929_1200"


def _write(path: Path, text: str):
    """精确落盘：不做换行转换（fixture 要控制 LF/CRLF 的真实字节）。"""
    with path.open("w", encoding="utf-8", newline="") as f:
        f.write(text)


@pytest.fixture()
def tree(tmp_path, monkeypatch):
    """最小产物树：L1 风险清单（1⬜+1✅）、L2 需求理解、L3 测试点、L4 竖排用例（CRLF）。"""
    out = tmp_path / "output"
    for rel in ("01_business", "02_split_requirements/知识管理",
                "03_design/test_points/知识管理", "03_design/test_cases/知识管理",
                "04_review", "05_export"):
        (out / rel).mkdir(parents=True)

    for attr, rel in [("BUSINESS_DIR", "01_business"), ("SPLIT_DIR", "02_split_requirements"),
                      ("TESTPOINTS_DIR", "03_design/test_points"),
                      ("CASES_DIR", "03_design/test_cases"),
                      ("REVIEW_DIR", "04_review"), ("EXPORT_DIR", "05_export"),
                      ("OUTPUT", ".")]:
        monkeypatch.setattr(lint, attr, out / rel)

    l1 = out / "02_split_requirements" / "知识管理" / "风险清单与待确认汇总.md"
    _write(l1, "\n".join([
        "# 风险清单与待确认汇总",
        "",
        "| 编号 | 类型 | 描述 | 确认结论(业务填) |",
        "|---|---|---|---|",
        "| R01 | 待确认 | 描述一 | ⬜ 问题一：是否级联删除？ |",
        "| R02 | 已澄清 | 描述二 | ✅ 已有结论 |",
    ]) + "\n")

    _write(out / "02_split_requirements" / "知识管理" / "01-分类管理.md",
           "# 分类管理\n\n| 规则 | 内容 |\n|---|---|\n| BR-01 | 分类可删除 |\n")

    l2 = out / "03_design" / "test_points" / "知识管理" / f"01-分类管理-需求理解_{_TS}.md"
    _write(l2, "\n".join([
        "# 分类管理 需求理解",
        "",
        "| 编号 | 问题 | 影响 | 确认结论(业务填) |",
        "|---|---|---|---|",
        "| Q01 | 删除规则 | 用例 | ⬜ 级联删除规则是什么？ |",
    ]) + "\n")

    l3 = out / "03_design" / "test_points" / "知识管理" / f"01-分类管理-测试点_{_TS}.md"
    _write(l3, "\n".join([
        "### 测试点表",
        "",
        "| 测试点编号 | 描述 | 是否依赖确认 | 确认结论(业务填) |",
        "|---|---|---|---|",
        "| TP-知识管理-分类管理-001 | 删除分类 | 是 | ⬜ 依赖 Q01：级联删除规则是什么？ |",
    ]) + "\n")

    l4 = out / "03_design" / "test_cases" / "知识管理" / f"01-分类管理-用例_{_TS}.md"
    _write(l4, "\r\n".join([
        "### TC-知识管理-分类管理-001 验证删除分类",
        "",
        "| 字段 | 内容 |",
        "|---|---|",
        "| 用例名称 | 删除分类 |",
        "| 确认结论(业务填) | ⬜ 级联删除规则是什么？ |",
        "",
    ]))

    return {"out": out, "l1": l1, "l2": l2, "l3": l3, "l4": l4, "export_dir": out / "05_export"}


def _exported_xlsx(tree) -> Path:
    xs = list(tree["export_dir"].glob("待拍板清单_*.xlsx"))
    assert len(xs) == 1, "export 应生成且仅生成一份待拍板清单"
    return xs[0]


def _fill(tree, xlsx, row_no: int, conclusion: str):
    wb = load_workbook(xlsx)
    ws = wb["待拍板"]
    ws.cell(row=row_no, column=6, value=conclusion)     # F 列 = 确认结论(业务填)
    wb.save(xlsx)


def test_export_collects_pending_rows_only(tree, capsys):
    """导出：收集 4 条 ⬜（横排 3 + 竖排 1），排除 ✅ 行；附填写说明 sheet。"""
    assert bridge.cmd_export([]) == 0                    # 无参数 -> 自动发现模块

    xlsx = _exported_xlsx(tree)
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["待拍板"]
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    assert len(rows) == 4
    questions = {r[4] for r in rows}
    assert "问题一：是否级联删除？" in questions
    assert "已有结论" not in str(questions)              # ✅ 行不进清单
    assert all(not r[5] for r in rows)                   # 结论列初始为空
    # 竖排字段行的编号/上下文退化为最近标题（用例 ID）
    l4_rows = [r for r in rows if r[2] == "用例"]
    assert l4_rows and l4_rows[0][3].startswith("TC-知识管理-分类管理-001")
    assert "填写说明" in wb.sheetnames


def test_apply_writes_back_in_place(tree):
    """回写：原位 ⬜->✅（单元格边界保留）、CRLF 不破坏、业务自带 ✅ 前缀不重复。"""
    assert bridge.cmd_export(["知识管理"]) == 0
    xlsx = _exported_xlsx(tree)
    _fill(tree, xlsx, 2, "级联删除，需先删子分类")          # L1 R01
    _fill(tree, xlsx, 5, "✅ 已带前缀结论")                 # L4 竖排行

    assert bridge.cmd_apply([str(xlsx)]) == 0

    l1 = tree["l1"].read_text(encoding="utf-8")
    assert "| R01 | 待确认 | 描述一 | ✅ 级联删除，需先删子分类 |" in l1
    assert "⬜ 问题一" not in l1
    assert "| R02 | 已澄清 | 描述二 | ✅ 已有结论 |" in l1   # 其他行原样

    l4 = tree["l4"].read_bytes().decode("utf-8")   # 原始字节读回：不被 universal newlines 归一
    assert "| 确认结论(业务填) | ✅ 已带前缀结论 |" in l4
    assert "✅ ✅" not in l4                               # 不重复加前缀
    assert "\r\n" in l4                                   # CRLF 文件换行风格逐字节保留
    assert "\r\n" not in tree["l1"].read_bytes().decode("utf-8")  # LF 文件不会被改写成 CRLF

    for key in ("l2", "l3"):                              # 未填行不动
        assert "⬜" in tree[key].read_text(encoding="utf-8")


def test_apply_stale_line_protection(tree, capsys):
    """过期保护：导出后源 md 行漂移 -> 跳过不误写，退出码 1。"""
    assert bridge.cmd_export(["知识管理"]) == 0
    xlsx = _exported_xlsx(tree)
    _fill(tree, xlsx, 2, "某结论")

    lines = tree["l1"].read_text(encoding="utf-8").splitlines(keepends=True)
    lines.insert(0, "（插入的新行，使后续行号整体 +1）\n")
    _write(tree["l1"], "".join(lines))

    assert bridge.cmd_apply([str(xlsx)]) == 1
    l1 = tree["l1"].read_text(encoding="utf-8")
    assert "⬜ 问题一：是否级联删除？" in l1               # 原行未被误写
    assert "✅ 某结论" not in l1
    assert "过期" in capsys.readouterr().out


def test_apply_dry_run_writes_nothing(tree):
    """--dry-run：走完整校验与预演报告，但不落盘。"""
    assert bridge.cmd_export(["知识管理"]) == 0
    xlsx = _exported_xlsx(tree)
    _fill(tree, xlsx, 2, "预演结论")

    assert bridge.cmd_apply([str(xlsx), "--dry-run"]) == 0
    assert "⬜ 问题一：是否级联删除？" in tree["l1"].read_text(encoding="utf-8")


def test_apply_without_fill_is_noop(tree, capsys):
    """业务一个都没填：无可回写、退出码 0、源 md 不动。"""
    assert bridge.cmd_export(["知识管理"]) == 0
    xlsx = _exported_xlsx(tree)

    assert bridge.cmd_apply([str(xlsx)]) == 0
    assert "未填写" in capsys.readouterr().out
    assert "⬜" in tree["l1"].read_text(encoding="utf-8")
