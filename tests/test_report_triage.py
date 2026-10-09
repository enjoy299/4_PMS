"""分流结论进报告：8 段摘要、有效通过率、BLOCKED 独立表、发版建议、旧路径不变。"""
import json

import pytest
from openpyxl import load_workbook

from core import config
from core.triage import append_records, new_record
from exporters.report_exporter import ReportExporter


def _rec(rid, cid, **kw):
    r = new_record(run_id=rid, case_id=cid, module="知识管理")
    r["name"] = kw.pop("name", f"用例{cid}")
    r["status"] = kw.pop("status", "fail")
    tri = kw.pop("triage", None)
    if tri is not None:
        r["triage"].update(tri)
    r.update(kw)
    return r


@pytest.fixture
def env(tmp_path, monkeypatch):
    """隔离目录与格式：不依赖真实 `output/`。"""
    monkeypatch.setattr(config, "TRIAGE_DIR", tmp_path / "triage")
    monkeypatch.setattr(config, "MD_REPORTS_DIR", tmp_path / "md")
    monkeypatch.setattr(config, "HTML_REPORTS_DIR", tmp_path / "html")
    monkeypatch.setattr(config, "EXCEL_REPORTS_DIR", tmp_path / "xlsx")
    monkeypatch.setattr(config, "REPORT_FORMAT", ["md"])
    monkeypatch.setattr(config, "HUMAN_CONFIRM", True)
    return tmp_path


def _seed(rid="R1"):
    """4 条覆盖四类终态：OK / ENV_BLOCK / DATA_PRECOND / R7→NEEDS_HUMAN。"""
    append_records([
        _rec(rid, "TC-001", status="pass",
             steps=[{"order": 1, "action": "点击新建", "outcome": "ok", "detail": ""}],
             assertions=[{"text": "弹窗关闭", "verdict": "pass", "reason": "", "evidence_refs": []}]),
        _rec(rid, "TC-002", status="error", error="net::ERR_CONNECTION_REFUSED",
             triage={"category": "ENV_BLOCK", "source": "rule", "confidence": 0.9,
                     "unblock_hint": "起本地服务后整 run 重跑"}),
        _rec(rid, "TC-003", status="fail", precondition_met=False,
             test_data="上传文件：test.doc",
             triage={"category": "DATA_PRECOND", "source": "rule", "confidence": 0.9,
                     "unblock_hint": "准备 test.doc 后重跑"}),
        _rec(rid, "TC-004", status="fail",
             assertions=[{"text": "分类树出现新分类", "verdict": "fail", "reason": "",
                          "evidence_refs": []}]),
    ], rid)
    return rid


def _read_md(tmp_path):
    files = sorted((tmp_path / "md").rglob("*.md"))
    assert files, "未生成 Markdown 报告"
    return files[-1].read_text(encoding="utf-8")


def test_summary_counts_and_effective_rate(env):
    report = ReportExporter().build_from_records(_seed(), mode="strict")
    s = report["summary"]
    assert (s["总用例"], s["通过"], s["环境阻塞"], s["数据前置未满足"], s["待人工裁决"]) == (4, 1, 1, 1, 1)
    assert s["移出分母"] == 2 and s["其他异常"] == 0
    assert s["有效通过率"] == round(1 / 2 * 100, 1)      # 1 / (4 - 1 - 1)
    assert s["通过率"] == round(1 / 4 * 100, 1)          # 旧口径仍保留、并列展示
    buckets = ("通过", "真实缺陷", "环境阻塞", "数据前置未满足", "用例问题",
               "执行未到位", "待人工裁决", "其他异常")
    assert sum(s[k] for k in buckets) == s["总用例"]      # 自检不变量


def test_md_has_effective_rate_blocked_table_and_triage_columns(env):
    _seed()
    ReportExporter().build_from_records("R1", mode="strict")
    md = _read_md(env)
    assert "## 摘要（8 段）" in md and "**有效通过率**" in md
    assert "| 分流类别 | 判定来源 | 置信度 | 裁决人 |" in md
    assert "## BLOCKED" in md and "起本地服务后整 run 重跑" in md    # 解除方式逐条列明
    assert "待人工裁决" in md and "TC-004" in md
    assert "分类树出现新分类" in md                                  # 断言原文仍可查


def test_strict_gate_rejects_release_advice_when_pending(env):
    _seed()
    report = ReportExporter().build_from_records("R1", mode="strict")
    assert report["gate"]["verdict"] == "reject" and report["gate"]["blocks_exit_code"]
    assert "不具备出结论条件" in _read_md(env)


def test_report_only_mode_writes_na_not_reject(env):
    _seed()
    report = ReportExporter().build_from_records("R1", mode="report_only")
    assert report["gate"]["verdict"] == "na" and not report["gate"]["blocks_exit_code"]
    md = _read_md(env)
    assert "N/A" in md and "浅分流" in md and "❌" not in md


def test_mode_defaults_from_human_confirm(env, monkeypatch):
    _seed()
    exporter = ReportExporter()
    assert exporter.build_from_records("R1")["mode"] == "strict"       # HUMAN_CONFIRM=True
    monkeypatch.setattr(config, "HUMAN_CONFIRM", False)
    assert exporter.build_from_records("R1")["mode"] == "report_only"  # CI 语义自动推得


def test_terminal_summary_printed_unless_quiet(env, capsys):
    """脚本直接调本层时给一屏摘要；被 `triage_pipeline` 驱动时静默，终端由编排层负责。"""
    _seed()
    ReportExporter().build_from_records("R1", mode="strict")
    assert "发版建议" in capsys.readouterr().out
    ReportExporter().build_from_records("R1", mode="strict", ts_file="T2", quiet=True)
    assert "发版建议" not in capsys.readouterr().out


def test_unknown_category_lands_in_other_bucket_not_dropped(env):
    """非法类别直写 jsonl（绕过 `append_records` 校验）时必须现形，不能静默消失。"""
    d = config.TRIAGE_DIR / "R9"
    d.mkdir(parents=True)
    rec = _rec("R9", "TC-901", status="fail", triage={"category": "WEIRD", "source": "rule"})
    (d / "records.jsonl").write_text(json.dumps(rec, ensure_ascii=False) + "\n", encoding="utf-8")
    s = ReportExporter().build_from_records("R9", mode="strict")["summary"]
    assert s["总用例"] == 1 and s["其他异常"] == 1
    assert "⚠️" in _read_md(env)


def test_blocked_table_follows_triage_category_not_raw_status(env):
    """BLOCKED 表必须与摘要的类别桶同一口径（真跑历史 run 时踩到的）。

    历史数据里原始 `status=blocked` 的用例远多于真正判为阻塞的：它们因未说明阻塞原因
    落入 NEEDS_HUMAN。若表按原始 status 取数，就会出现「摘要 5 条阻塞、表里 11 行」，
    读者不知道哪些才是真的逃出了有效分母。
    """
    append_records([
        _rec("RB", "TC-B1", status="blocked",
             triage={"category": "ENV_BLOCK", "source": "rule", "confidence": 0.9,
                     "unblock_hint": "起本地服务后重跑"}),
        _rec("RB", "TC-B2", status="blocked"),      # 只写了 blocked、未说明原因 → R8 送人工
    ], "RB")
    report = ReportExporter().build_from_records("RB", mode="strict")
    assert report["summary"]["环境阻塞"] == 1 and report["summary"]["待人工裁决"] == 1
    assert [d["id"] for d in report["blocked"]] == ["TC-B1"]
    md = _read_md(env)
    table = md.split("## BLOCKED")[1].split("## 待人工裁决")[0]
    assert "TC-B1" in table and "TC-B2" not in table
    # 结论列也不能自相矛盾：一条同时写 BLOCKED 与 待人工裁决 的行，人不知道先处理哪个
    assert "### [待人工裁决] [TC-B2]" in md


def test_suite_name_is_majority_module_not_first(env):
    """报告标题取多数模块（真跑历史 run 时踩到的）。

    `P0_km_20260907` 的 12 条里 11 条属知识管理、1 条属分类管理，而分类管理那条恰好
    Case ID 排序靠首。取首条会让整份报告改名，人拿文件时会误以为拿错了模块。
    """
    append_records([
        _rec("RM", "TC-CAT-001", module="分类管理"),
        _rec("RM", "TC-KM-001", module="知识管理"),
        _rec("RM", "TC-KM-002", module="知识管理"),
    ], "RM")
    report = ReportExporter().build_from_records("RM", mode="report_only")
    assert report["suite_name"] == "知识管理"


def test_blank_case_name_is_marked_not_left_empty(env):
    """历史 JSON 的 `pending_cases` 只存 case_id + blocker，名称必为空。

    空单元格会让人以为是报告渲染坏了（或直接误读成「这条用例没名字」），必须显式标出来。
    """
    append_records([_rec("RN", "TC-N01", name="", status="blocked")], "RN")
    ReportExporter().build_from_records("RN", mode="strict")
    assert "（名称未记录" in _read_md(env)


def test_missing_run_raises_clearly(env):
    with pytest.raises(FileNotFoundError) as e:
        ReportExporter().build_from_records("nope")
    assert "records.jsonl" in str(e.value)


@pytest.fixture
def env_web(env, monkeypatch):
    """复用 `env`，只把输出格式换成 HTML + Excel。"""
    monkeypatch.setattr(config, "REPORT_FORMAT", ["html", "excel"])
    return env


def _read_html(tmp_path):
    files = sorted((tmp_path / "html").rglob("*.html"))
    assert files, "未生成 HTML 报告"
    return files[-1].read_text(encoding="utf-8")


def _load_xlsx(tmp_path):
    files = sorted((tmp_path / "xlsx").rglob("*.xlsx"))
    assert files, "未生成 Excel 报告"
    return load_workbook(files[-1])


def test_html_has_six_cards_and_effective_rate(env_web):
    _seed()
    ReportExporter().build_from_records("R1", mode="strict")
    html = _read_html(env_web)
    assert html.count('<div class="card ') == 6
    for label in ("Total", "Passed", "Product Bugs", "Env / Data Blocked", "Pending Review",
                  "Effective Pass Rate"):
        assert label in html
    assert ">50.0%<" in html                      # 有效通过率（不是原始 25.0%）
    assert "发版建议" in html and "不具备出结论条件" in html


def test_html_blocked_section_and_triage_headers(env_web):
    _seed()
    ReportExporter().build_from_records("R1", mode="strict")
    html = _read_html(env_web)
    assert "BLOCKED（环境 / 数据前置，单独成表）" in html
    assert "起本地服务后整 run 重跑" in html
    detail_head = html.split("<table>")[1].split("</tr>")[0]   # 只数明细表的列头，不含 BLOCKED 表
    assert detail_head.count("<th>") == 18                       # 14 旧列 + 4 分流列
    for h in ("分流类别", "判定来源", "置信度", "裁决人"):
        assert f"<th>{h}</th>" in html


def test_html_escapes_angle_brackets(env_web):
    """断言原文含 `<b>` 时不得破坏表格（旧版未转义，本任务一并修）。"""
    append_records([_rec("R8", "TC-050", status="fail",
                         assertions=[{"text": "提示「<b>保存成功</b>」出现", "verdict": "fail",
                                      "reason": "", "evidence_refs": []}])], "R8")
    ReportExporter().build_from_records("R8", mode="strict")
    html = _read_html(env_web)
    assert "&lt;b&gt;保存成功&lt;/b&gt;" in html


def test_excel_has_18_columns_and_blocked_sheet(env_web):
    _seed()
    ReportExporter().build_from_records("R1", mode="strict")
    wb = _load_xlsx(env_web)
    ws = wb["Test Results"]
    heads = [c.value for c in ws[4]]
    assert len(heads) == 18 and heads[-4:] == ["分流类别", "判定来源", "置信度", "裁决人"]
    assert "BLOCKED" in wb.sheetnames
    blocked = [[c.value for c in row] for row in wb["BLOCKED"].iter_rows(min_row=2)]
    assert any(r[0] == "TC-002" and "起本地服务" in str(r[5]) for r in blocked)
    assert any(r[0] == "TC-003" for r in blocked)
    assert len(blocked) == 2


def test_excel_conclusion_uses_blocked_and_pending(env_web):
    _seed()
    ReportExporter().build_from_records("R1", mode="strict")
    ws = _load_xlsx(env_web)["Test Results"]
    by_id = {row[0].value: row[5].value for row in ws.iter_rows(min_row=5)}
    assert by_id["TC-001"] == "PASSED"
    assert by_id["TC-002"] == "BLOCKED" and by_id["TC-003"] == "BLOCKED"
    assert by_id["TC-004"] == "待人工裁决"
    assert any(row[14].value == "环境阻塞" for row in ws.iter_rows(min_row=5))   # O 列分流类别


def test_html_conclusion_color_follows_category(env_web):
    """结论单元格按分流类别着色，不跟着原始 `status` 走。

    原始颜色由 `status-fail` 提供，于是「待人工裁决」被刷成缺陷红——未定论的用例看起来
    像已确认失败，正好把分流最想传达的信息说反。
    """
    _seed()
    ReportExporter().build_from_records("R1", mode="strict")
    html = _read_html(env_web)
    assert 'class="c-pending">待人工裁决' in html
    assert 'class="c-blocked">BLOCKED' in html
    assert 'class="c-ok">PASSED' in html


def test_excel_status_fill_follows_category(env_web):
    """Excel 的 Status 列底色同理：分流模式下跟类别，不跟原始四态。"""
    _seed()
    ReportExporter().build_from_records("R1", mode="strict")
    ws = _load_xlsx(env_web)["Test Results"]
    fill = {row[0].value: str(row[5].fill.start_color.rgb) for row in ws.iter_rows(min_row=5)}
    assert fill["TC-004"] == "00FFF9C4"       # 待裁决黄
    assert fill["TC-002"] == "00E1BEE7"       # 环境阻塞紫
    assert fill["TC-001"] == "00E2EFDA"       # 通过绿


def test_blank_name_marked_in_html_and_excel(env_web):
    """历史 run 的阻塞用例没存名称：HTML/Excel 明细表也得显式标出来。

    MD 侧 Task 11 已处理；这里若留空，整列空白看起来像渲染坏了而不像「数据本就没有」。
    """
    append_records([_rec("R9", "TC-060", name="", status="blocked")], "R9")
    ReportExporter().build_from_records("R9", mode="strict")
    assert "（名称未记录" in _read_html(env_web)
    ws = _load_xlsx(env_web)["Test Results"]
    assert [row[1].value for row in ws.iter_rows(min_row=5)] == ["（名称未记录，见用例库）"]


def test_legacy_html_and_excel_shape_unchanged(env_web):
    """旧路径：14 列 / 5 卡片，一行结论不得变。"""
    rows = [{"id": "TC-001", "name": "n", "module": "m", "case_type": "正向", "priority": "P0",
             "status": "pass", "duration": 1.0, "steps": "1. a", "expected": "1. b",
             "assertions": [{"text": "t", "passed": True}]}]
    ReportExporter().build_and_save("知识管理", rows)
    html = _read_html(env_web)
    assert html.count('<div class="card ') == 5 and html.count("<th>") == 14
    wb = _load_xlsx(env_web)
    assert len([c.value for c in wb["Test Results"][4]]) == 14
    assert "BLOCKED" not in wb.sheetnames


def test_legacy_build_and_save_still_four_state(env):
    """旧路径一行结论不得变：`TRIAGE_ENABLED=false` 时走的就是它。"""
    rows = [{"id": "TC-001", "name": "n", "module": "m", "case_type": "正向", "priority": "P0",
             "status": "pass", "duration": 1.0, "steps": "1. a", "expected": "1. b",
             "assertions": [{"text": "t", "passed": True}]}]
    ReportExporter().build_and_save("知识管理", rows)
    md = _read_md(env)
    assert "| Passed | 1 |" in md and "分流类别" not in md and "摘要（8 段）" not in md
